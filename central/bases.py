"""Carrega as bases mais recentes da fonte e as deixa prontas para uso
(colunas padronizadas, datas e valores convertidos, campos calculados).

Tudo é cacheado pela "assinatura" do arquivo (nome + data de modificação +
tamanho): enquanto o arquivo não muda, ninguém relê; quando chega um arquivo
novo, a próxima execução já usa ele — sem botão de importar, sem recarregar.
"""

from __future__ import annotations

import io
import json
import re
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import PurePath

import numpy as np
import pandas as pd
import streamlit as st

from . import af as af_mod
from . import config, fontes, fotos, leitura
from .leitura import ACAO_AF, AF, CONF, EQUIP, EQUIPE, IP19, IW38, MB52, NOTAS, OPER, REQ, TIPOS
from .util import achar_coluna, agora_local, chave, fora_da_visao, hoje_local, marcar_quebras, para_data, para_numero, sem_acento, setores, texto

ARQ_CAD_EQUIP = "cadastro_equipamentos.json"
ARQ_CAD_MAT = "cadastro_materiais.json"
ARQ_PREF = "preferencias.json"

# Pista pelo nome do arquivo: arquivos com pista são verificados primeiro.
_PISTAS = [(AF, r"GERENCIADOR DE AF|ANALISE.*FALHA|\bAF\b"), (CONF, r"IW47|IW41|CONFIRMA|APONTAMENTO"),
           (EQUIPE, r"GESTAO.*HH|EQUIPE|EFETIVO"),
           (OPER, r"IW38OP|IW37|OPERAC"), (IP19, r"IP19|IP24|PLANOS?\b"), (IW38, r"IW38|IW39|ORDENS?"),
           (NOTAS, r"IW28|IW29|NOTAS?\b"), (EQUIP, r"IH08|IE05|EQUIPAMENT"), (MB52, r"MB52|MB51|ESTOQUE|MATERIA"),
           (REQ, r"REQUISI|REQ\b|APROVA|COMPRAS|SOLICITA")]


# ----------------------------------------------------------------------------
# Fonte e inventário de arquivos
# ----------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def _fonte_cacheada(chave_cfg: str, versao_modulo: int):
    return fontes.criar(config.carregar())


def fonte():
    # versao_modulo: quando o código é atualizado (nova publicação na nuvem) o módulo fontes
    # é recarregado; a fonte antiga criaria Arquivo da classe velha, que o cache não grava.
    return _fonte_cacheada(config.carregar().chave, id(fontes.Arquivo))


@st.cache_data(show_spinner=False, max_entries=4)
def _listar(chave_cfg: str, janela: int) -> list[fontes.Arquivo]:  # janela = chave de tempo do cache
    lista = sorted(fonte().listar(), key=lambda a: a.modificado, reverse=True)
    # Sempre com a classe atual (ver fonte()), senão o cache falha com UnserializableReturnValueError.
    return [fontes.Arquivo(a.id, a.nome, a.modificado, a.tamanho, a.versao) for a in lista]


def arquivos() -> list[fontes.Arquivo]:
    """Lista da fonte, compartilhada por todos os usuários. Pastas locais são relidas a
    cada 15 s; o SharePoint, uma vez por intervalo de verificação (poupa a API)."""
    cfg = config.carregar()
    passo = cfg.intervalo_verificacao if cfg.fonte == "sharepoint" else 15
    return _listar(cfg.chave, int(time.time() // passo))


@st.cache_data(show_spinner=False, max_entries=8, ttl="30m")
def _bytes(id: str, assinatura: str) -> bytes:
    return fonte().ler(id)


@st.cache_data(show_spinner=False, max_entries=6, ttl="2h")
def _ler_cru(id: str, assinatura: str, planilha: str | None, linha: int | None) -> pd.DataFrame:
    return leitura.ler_arquivo(id, _bytes(id, assinatura), planilha, linha)


@st.cache_data(show_spinner=False, max_entries=2048)
def _sondar(id: str, assinatura: str) -> list[tuple[str, str | None, int | None]]:
    """Bases dentro do arquivo. Falha de leitura (arquivo aberto no Excel, OneDrive sincronizando, rede)
    levanta erro e não fica em cache: o arquivo é tentado de novo, em vez de sumir até mudar."""
    try:
        return leitura.sondar(id, fonte().ler(id))
    except leitura.LeituraErro:  # não é uma base que o app conhece
        return []


_FALHAS_SONDA: dict[tuple[str, str], float] = {}   # (arquivo, versão) → hora da última falha de leitura


def _sondar_seguro(arq: fontes.Arquivo) -> tuple[list, bool]:
    """(bases achadas, leu?) — depois de uma falha, espera 1 min antes de tentar o mesmo arquivo de novo."""
    chave_f = (arq.id, arq.assinatura)
    falhou = _FALHAS_SONDA.get(chave_f)
    if falhou and time.time() - falhou < 60:
        return [], False
    try:
        achados = _sondar(arq.id, arq.assinatura)
    except Exception:  # noqa: BLE001
        _FALHAS_SONDA[chave_f] = time.time()
        return [], False
    _FALHAS_SONDA.pop(chave_f, None)
    return achados, True


@dataclass(frozen=True)
class Origem:
    arquivo: fontes.Arquivo
    tipo: str
    planilha: str | None
    linha: int | None

    @property
    def rotulo(self) -> str:
        return self.arquivo.nome + (f" › aba {self.planilha}" if self.planilha else "")


@dataclass
class Inventario:
    ativos: dict[str, Origem] = field(default_factory=dict)       # tipo -> arquivo principal (o mais novo)
    usados: dict[str, list[Origem]] = field(default_factory=dict)  # tipo -> todos os arquivos somados
    candidatos: dict[str, list[Origem]] = field(default_factory=dict)  # tipo -> todas encontradas
    fixados: dict[str, str] = field(default_factory=dict)          # tipo -> id fixado pelo usuário
    verificados: int = 0
    erro: str = ""
    incompleto: bool = False   # algum arquivo não pôde ser lido agora (não guardar este inventário)


def _pista(arq: fontes.Arquivo) -> str | None:
    base = chave(PurePath(arq.arquivo).stem)
    return next((t for t, rx in _PISTAS if re.search(rx, base)), None)


_INV_MEMO: dict[str, Inventario] = {}


def inventario() -> Inventario:
    """Inventário das bases; recalculado só quando algum arquivo da fonte muda."""
    try:
        chave_memo = config.carregar().chave + "#" + assinatura_geral()
    except Exception:  # noqa: BLE001 — fonte fora do ar: _inventario mostra o erro
        return _inventario()
    if chave_memo not in _INV_MEMO:
        inv = _inventario()
        if inv.incompleto or inv.erro:   # falha passageira: tenta de novo na próxima vez
            return inv
        _INV_MEMO.clear()
        _INV_MEMO[chave_memo] = inv
    return _INV_MEMO[chave_memo]


def _inventario() -> Inventario:
    """Percorre os arquivos do mais novo para o mais velho e identifica cada base.

    Arquivos com nome de export do SAP (IW38, IP19, MB52…) são sempre verificados;
    os demais só até todas as bases terem sido achadas — um arquivo qualquer mais
    velho que o mais antigo em uso não tem como ser "o mais recente". A
    identificação de cada arquivo fica em cache até ele mudar.
    """
    inv = Inventario()
    try:
        lista = arquivos()
    except Exception as e:  # noqa: BLE001
        inv.erro = str(e)
        return inv
    inv.fixados = {k: v for k, v in ler_cadastro(ARQ_PREF).items() if isinstance(v, str)}
    dados = [a for a in lista if not a.arquivo.lower().endswith(".json")]
    # arquivos com pista no nome primeiro (mantendo a ordem por data dentro de cada grupo)
    dados.sort(key=lambda a: _pista(a) is None)
    pendentes = set(TIPOS) - {t for t, id_ in inv.fixados.items() if any(a.id == id_ for a in dados)}
    limite = None  # data do arquivo mais velho em uso entre os tipos já achados
    fixos = set(inv.fixados.values())
    for arq in dados:
        if (not pendentes and arq.id not in fixos and _pista(arq) is None
                and (limite is None or arq.modificado < limite)):
            continue
        achados, leu = _sondar_seguro(arq)
        inv.incompleto |= not leu
        inv.verificados += 1
        for tipo, planilha, linha in achados:
            o = Origem(arq, tipo, planilha, linha)
            inv.candidatos.setdefault(tipo, []).append(o)
            if tipo in pendentes:
                pendentes.discard(tipo)
                limite = arq.modificado if limite is None else min(limite, arq.modificado)
    for tipo, lista_tipo in inv.candidatos.items():
        lista_tipo.sort(key=lambda o: o.arquivo.modificado, reverse=True)
        fixo = inv.fixados.get(tipo)
        escolhido = next((o for o in lista_tipo if o.arquivo.id == fixo), None) if fixo else None
        inv.usados[tipo] = [escolhido] if escolhido else leitura.escolher_origens(tipo, lista_tipo)
        inv.ativos[tipo] = inv.usados[tipo][0]
    return inv


def fixar_origem(tipo: str, arquivo_id: str | None, usuario: str = "") -> None:
    """Fixa um arquivo para uma base (ele continua sendo relido quando é sobrescrito)
    ou volta para o modo automático (arquivo_id=None)."""
    atual = _ler_json_agora(ARQ_PREF)
    if arquivo_id:
        atual[tipo] = arquivo_id
    else:
        atual.pop(tipo, None)
    fonte().gravar(ARQ_PREF, json.dumps(atual, ensure_ascii=False, indent=1).encode("utf-8"))
    recarregar()


def assinatura_geral() -> str:
    """Muda quando qualquer arquivo da fonte muda (chave do inventário)."""
    return ";".join(a.assinatura for a in arquivos())


# cadastros que entram nos indicadores (os de controle — usuários, lotes do robô, sinal do PC, backups —
# não recarregam a tela de ninguém nem invalidam os indicadores)
CADASTROS_DOS_DADOS = ("metas_wcm.json", ARQ_CAD_EQUIP, ARQ_CAD_MAT, ARQ_PREF)


def assinatura_dados() -> str:
    """Muda quando uma base do SAP em uso ou um cadastro que entra nos indicadores muda — é o que o
    vigia da barra lateral observa e a chave dos indicadores em cache."""
    inv = inventario()
    partes = [o.arquivo.assinatura for t in sorted(inv.usados) for o in inv.usados[t]]
    ids = {fonte().id_de(n) for n in CADASTROS_DOS_DADOS}
    partes += sorted(a.assinatura for a in arquivos() if a.id in ids)
    return ";".join(partes)


# ----------------------------------------------------------------------------
# IW38 — ordens de manutenção
# ----------------------------------------------------------------------------
_IW38_COLUNAS = {
    "Ordem": [r"^ORDEM$"],
    "Tipo": [r"TIPO DE ORDEM", r"^TIPO ORDEM"],
    "Prioridade": [r"^PRIORIDADE$", r"PRIORIDADE"],
    "Texto": [r"TEXTO BREVE", r"^DESCRICAO$"],
    "Equipamento": [r"^EQUIPAMENTO$"],
    "Objeto técnico": [r"DENOMINACAO DO OBJETO TECNICO", r"DENOMINACAO.*EQUIP", r"OBJETO TECNICO"],
    "Local de instalação": [r"DENOMINACAO DO LOC INSTALACAO", r"DENOMINACAO.*LOC", r"LOCAL DE INSTALACAO"],
    "Centro de trabalho": [r"CENTRO TRAB", r"CENTRO DE TRABALHO"],
    "Centro de custo": [r"CENTRO CUSTO", r"CENTRO DE CUSTO"],
    "Localização": [r"^LOCALIZACAO$"],
    "Centro": [r"^CENTRO$"],
    "Plano": [r"^PLANO DE MANUTENCAO$", r"^PLANO MANUT", r"^PLANO$"],
    "Status sistema": [r"STATUS DO SISTEMA", r"STATUS SISTEMA"],
    "Status usuário": [r"STATUS USUARIO"],
    "Início": [r"DATA BASE DO INICIO", r"INICIO BASE", r"DATA INICIO"],
    "Fim": [r"DATA BASE DO FIM", r"FIM BASE", r"DATA FIM"],
    "Entrada": [r"DATA DE ENTRADA", r"^CRIADO EM$"],
    "Criado por": [r"CRIADO POR"],
    "Custo real": [r"CUSTOS? TOT.*REA", r"CUSTO.*REA", r"CUSTOS? TOT"],
    "Custo planejado": [r"CUSTOS? TOT.*PLAN", r"CUSTO.*PLAN", r"CUSTO.*ESTIM"],
}

# Status de usuário conhecidos (os demais aparecem pelo código).
STATUS_USUARIO = {
    "PLA": "Em planejamento", "PRO": "Programado", "ELP": "Em planejamento (ELP)", "CAN": "Cancelada",
}

ENCERRADA_SEM_CONF = "Encerrada sem confirmação"
SITUACOES = ["Aberta", "Liberada", ENCERRADA_SEM_CONF, "Concluída", "Cancelada", "Sem status"]
NAT_PLANO, NAT_BACKLOG = "Plano de manutenção", "Backlog"
NATUREZAS = [NAT_PLANO, NAT_BACKLOG]
PENDENTES = ("Aberta", "Liberada", ENCERRADA_SEM_CONF)  # tudo que ainda não está concluído


def _situacao(sistema: str, usuario: str) -> str:
    sis = set(sistema.split())
    usu = set(usuario.split())
    if "CAN" in usu or sis & {"DLFL", "MREL", "MEEL"}:
        return "Cancelada"
    # regra do PCM: concluída só com CONF (confirmada) + ENTE (encerrada tecnicamente) juntos
    if {"CONF", "ENTE"} <= sis:
        return "Concluída"
    if sis & {"ENCE", "ENTE"}:
        return ENCERRADA_SEM_CONF
    if "LIB" in sis:
        return "Liberada"
    if "ABER" in sis:
        return "Aberta"
    return "Sem status"


def _piloto_ou_matriz(df: pd.DataFrame) -> pd.Series:
    """Fábrica Piloto, Desenho (localização DESEN) e OPER_MATRIZ ficam fora (ver util.fora_da_visao): usa centro
    de trabalho, localização, local de instalação e objeto — nunca o texto livre ("MESA DE FIX. MATRIZ" fica)."""
    return fora_da_visao(df)


def _cru(tipo: str, origens: tuple) -> pd.DataFrame:
    """Tabela crua da base (o export mais recente; ver leitura.MESCLAR). Sem cache próprio: só é usada para
    preparar a base (que fica em cache) — guardar a crua também dobraria a memória."""
    return leitura.mesclar(tipo, [_ler_cru(o.arquivo.id, o.arquivo.assinatura, o.planilha, o.linha) for o in origens])


@st.cache_resource(show_spinner="Preparando ordens (IW38)…", max_entries=2)
def _iw38(origens: tuple, excluir: bool, dia: str) -> tuple[pd.DataFrame, int]:
    """`dia` na chave: atrasada e dias em aberto dependem de hoje (o cache vira à meia-noite)."""
    return preparar_iw38(_cru(IW38, origens), pd.Timestamp(dia), excluir_piloto_matriz=excluir)


def preparar_iw38(cru: pd.DataFrame, hoje: pd.Timestamp | None = None,
                  excluir_piloto_matriz: bool = True) -> tuple[pd.DataFrame, int]:
    hoje = (hoje or agora_local()).normalize()
    df = pd.DataFrame(index=cru.index)
    usadas = set()
    for nome, padroes in _IW38_COLUNAS.items():
        col = achar_coluna([c for c in cru.columns if c not in usadas], *padroes,
                           excluir=r"TEXTO|ITEM" if nome == "Plano" else None)
        if col:
            usadas.add(col)
        df[nome] = cru[col] if col else ""

    for c in ["Ordem", "Tipo", "Prioridade", "Texto", "Equipamento", "Objeto técnico", "Local de instalação",
              "Centro de trabalho", "Centro de custo", "Localização", "Centro", "Plano", "Status sistema",
              "Status usuário", "Criado por"]:
        df[c] = texto(df[c])
    # linhas sem ordem (ex.: fórmula de "Status" arrastada até o fim da planilha) não são ordens
    df = df[df["Ordem"] != ""]
    cru = cru.loc[df.index]
    df["Status sistema"] = df["Status sistema"].str.split().str.join(" ")
    df["Status usuário"] = df["Status usuário"].str.split().str.join(" ")
    for c in ["Início", "Fim", "Entrada"]:
        df[c] = para_data(df[c])
    for c in ["Custo real", "Custo planejado"]:
        df[c] = para_numero(df[c]).fillna(0.0)

    # colunas extras do export (que o app não conhece) continuam disponíveis
    for c in cru.columns:
        if c not in usadas and c not in df.columns:
            df[c] = cru[c]

    fora = _piloto_ou_matriz(df) if excluir_piloto_matriz else pd.Series(False, index=df.index)
    removidas = int(fora.sum())
    df = df.loc[~fora].reset_index(drop=True)

    # regra do PCM: número na coluna Plano = ordem de plano de manutenção; sem número = backlog
    df["Com plano"] = df["Plano"].str.contains(r"\d", regex=True)
    df["Natureza"] = np.where(df["Com plano"], NAT_PLANO, NAT_BACKLOG)
    pares = pd.Series(list(zip(df["Status sistema"], df["Status usuário"])))
    cache = {p: _situacao(*p) for p in set(pares)}
    df["Situação"] = pd.Categorical(pares.map(cache), categories=SITUACOES)
    df["Data"] = df["Início"].fillna(df["Entrada"]).fillna(df["Fim"])
    pendente = df["Situação"].isin(PENDENTES)
    prazo = df["Fim"].fillna(df["Início"])
    df["Atrasada"] = pendente & prazo.notna() & (prazo < hoje)
    # sem data de entrada no export, conta do início previsto (ordem programada para o futuro: 0, não negativo)
    df["Dias em aberto"] = np.where(pendente, (hoje - df["Entrada"].fillna(df["Início"])).dt.days.clip(lower=0),
                                    np.nan)
    dur = (df["Fim"] - df["Início"]).dt.days
    df["Duração (dias)"] = dur.where((dur >= 0) & (dur < 3650))
    df["Equip. (chave)"] = np.where(df["Equipamento"] != "", df["Equipamento"], df["Objeto técnico"])
    busca = df["Ordem"]
    for c in ["Texto", "Equipamento", "Objeto técnico", "Local de instalação", "Plano", "Criado por", "Status usuário"]:
        busca = busca + " " + df[c]
    df["_busca"] = busca.map(lambda s: sem_acento(s).upper())
    df["Setor"] = setores(df["Centro de trabalho"])
    return df, removidas


# ----------------------------------------------------------------------------
# MB52 — estoque
# ----------------------------------------------------------------------------
@st.cache_resource(show_spinner="Preparando estoque (MB52)…", max_entries=2)
def _mb52(origens: tuple) -> tuple[pd.DataFrame, pd.DataFrame]:
    return preparar_mb52(_cru(MB52, origens))


def preparar_mb52(cru: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Devolve (resumo por material, detalhe por depósito)."""
    cols = list(cru.columns)
    mat = achar_coluna(cols, r"^MATERIAL$", r"CODIGO", r"MATERIAL", excluir=r"TEXTO|DESCR|TIPO|GRUPO")
    if not mat:
        raise leitura.LeituraErro("coluna de código do material não encontrada no MB52")
    desc = achar_coluna(cols, r"TEXTO BREVE", r"DESCRICAO", r"TEXTO")
    dep = achar_coluna(cols, r"^DEPOSITO$", r"DEPOSITO", excluir=r"DENOMINACAO|DESCR")
    um = achar_coluna(cols, r"UNIDADE MEDIDA", r"^UMB$", r"^UM$", r"UNID")
    qtd = achar_coluna(cols, r"UTILIZACAO LIVRE", r"UTILIZ LIVRE", r"LIVRE UTILIZ", r"^ESTOQUE$", r"ESTOQUE",
                       r"QUANTIDADE", r"DISPONIVEL", excluir=r"VALOR|MOEDA")
    val = achar_coluna(cols, r"VALOR UTILIZ", r"VALOR.*LIVRE", r"^VALOR$", r"VALOR")

    det = pd.DataFrame({
        "Material": texto(cru[mat]),
        "Descrição": texto(cru[desc]) if desc else "",
        "Depósito": texto(cru[dep]) if dep else "",
        "UM": texto(cru[um]) if um else "",
        "Estoque": para_numero(cru[qtd]).fillna(0.0) if qtd else 0.0,
        "Valor": para_numero(cru[val]).fillna(0.0) if val else 0.0,
    })
    det = det[det["Material"] != ""]
    resumo = det.groupby("Material", as_index=False, sort=False).agg(
        Descrição=("Descrição", "first"),
        UM=("UM", "first"),
        Estoque=("Estoque", "sum"),
        Valor=("Valor", "sum"),
        Depósitos=("Depósito", lambda s: ", ".join(sorted({d for d in s if d}))),
    )
    return resumo, det.reset_index(drop=True)


# ----------------------------------------------------------------------------
# Requisições de compra (aprovações)
# ----------------------------------------------------------------------------
_REQ_COLUNAS = {
    "Requisição": [r"N(UMERO)? DA REQ", r"REQUISI"],
    "PO": [r"NUMERO DA PO", r"PEDIDO", r"\bPO\b"],
    "Status": [r"^STATUS$", r"STATUS"],
    "Aprovador": [r"APROVADOR ATUAL", r"APROVADOR"],
    "Proprietário da aprovação": [r"PROPRIETARIO"],
    "Solicitante": [r"SOLICITADO POR", r"SOLICITANTE", r"REQUISITANTE"],
    "Itens": [r"^ITENS$", r"ITENS", r"DESCRICAO"],
    "Justificativa": [r"JUSTIFICATIVA"],
    "Enviado em": [r"^ENVIADO EM$", r"ENVIADO"],
    "Último envio": [r"ULTIMO ENVIO"],
    "Total": [r"^TOTAL$", r"VALOR TOTAL", r"TOTAL"],
}


@st.cache_resource(show_spinner="Preparando requisições…", max_entries=2)
def _req(origens: tuple, dia: str) -> pd.DataFrame:
    return preparar_requisicoes(_cru(REQ, origens), pd.Timestamp(dia))


def preparar_requisicoes(cru: pd.DataFrame, hoje: pd.Timestamp | None = None) -> pd.DataFrame:
    hoje = (hoje or agora_local()).normalize()
    df = pd.DataFrame(index=cru.index)
    usadas = set()
    for nome, padroes in _REQ_COLUNAS.items():
        col = achar_coluna([c for c in cru.columns if c not in usadas], *padroes)
        if col:
            usadas.add(col)
        df[nome] = cru[col] if col else ""
    for c in ["Requisição", "PO", "Status", "Aprovador", "Proprietário da aprovação", "Solicitante", "Itens", "Justificativa"]:
        df[c] = texto(df[c])
    df["Enviado em"] = para_data(df["Enviado em"])
    df["Último envio"] = para_data(df["Último envio"])
    df["Total"] = para_numero(df["Total"]).fillna(0.0)
    df["Pendente"] = df["Status"].map(chave).str.contains("PENDENTE")
    ref = df["Último envio"].fillna(df["Enviado em"])
    df["Dias aguardando"] = np.where(df["Pendente"], (hoje - ref).dt.days, np.nan)
    return df


# ----------------------------------------------------------------------------
# IP19 — programação dos planos de manutenção
# ----------------------------------------------------------------------------
@st.cache_resource(show_spinner="Preparando planos (IP19)…", max_entries=2)
def _ip19(origens: tuple, excluir: bool) -> pd.DataFrame:
    from .planos import preparar_ip19

    return preparar_ip19(_cru(IP19, origens), excluir_piloto_matriz=excluir)


# ----------------------------------------------------------------------------
# IW38OP — operações das ordens (trabalho planejado em horas)
# ----------------------------------------------------------------------------
_OPER_COLUNAS = {
    "Ordem": [r"^ORDEM$"], "Operação": [r"^OPERACAO$"], "Tipo": [r"TIPO DE ORDEM"],
    "Centro de trabalho": [r"^CENTRO DE TRABALHO$", r"CENTRO TRAB"],
    "Equipamento": [r"^EQUIPAMENTO$"], "Objeto técnico": [r"DENOMINACAO DO OBJETO", r"OBJETO TECNICO"],
    "Local de instalação": [r"DENOMINACAO DO LOC"], "Texto da operação": [r"TXT BREVE OPERACAO", r"TEXTO.*OPERACAO"],
    "Trabalho": [r"^TRABALHO$"], "Unidade": [r"UNIDADE DO TRABALHO", r"UNID.*TRAB"],
    "Pessoas": [r"^NUMERO$", r"N PESSOAS"], "Status sistema": [r"STATUS DO SISTEMA"],
    "Início": [r"1A DATA DE INICIO", r"DATA DE INICIO", r"ULTIMA DATA INICIO"], "Fim real": [r"DATA DO FIM REAL"],
}
_HORAS_POR = {"MIN": 1 / 60, "H": 1, "HR": 1, "STD": 1, "HRS": 1, "D": 8, "DIA": 8}


def _padronizar(cru: pd.DataFrame, mapa: dict, textos: list[str]) -> pd.DataFrame:
    df = pd.DataFrame(index=cru.index)
    usadas: set[str] = set()
    for nome, padroes in mapa.items():
        col = achar_coluna([c for c in cru.columns if c not in usadas], *padroes)
        if col:
            usadas.add(col)
        df[nome] = cru[col] if col else ""
    for c in textos:
        df[c] = texto(df[c])
    return df


def preparar_operacoes(cru: pd.DataFrame, excluir_piloto_matriz: bool = True) -> pd.DataFrame:
    df = _padronizar(cru, _OPER_COLUNAS, ["Ordem", "Operação", "Tipo", "Centro de trabalho", "Equipamento",
                                           "Objeto técnico", "Local de instalação", "Texto da operação", "Unidade",
                                           "Status sistema"])
    df = df[df["Ordem"] != ""].copy()
    if excluir_piloto_matriz:
        df = df[~fora_da_visao(df)].copy()
    fator = df["Unidade"].str.upper().map(_HORAS_POR).fillna(1 / 60)
    df["Horas"] = (para_numero(df["Trabalho"]).fillna(0.0) * fator).round(2)
    df["Pessoas"] = para_numero(df["Pessoas"]).fillna(0).astype(int)
    for c in ["Início", "Fim real"]:
        df[c] = para_data(df[c])
    toks = df["Status sistema"].str.split().map(set)
    df["Concluída"] = toks.map(lambda t: bool(t & {"CONF", "ENTE", "ENCE"}))
    df["Cancelada"] = toks.map(lambda t: bool(t & {"DLFL", "MREL", "MEEL"}))
    df["Setor"] = setores(df["Centro de trabalho"])
    return df.drop(columns=["Trabalho"]).reset_index(drop=True)


@st.cache_resource(show_spinner="Preparando operações (IW38OP)…", max_entries=2)
def _oper(origens: tuple, excluir: bool) -> pd.DataFrame:
    return preparar_operacoes(_cru(OPER, origens), excluir)


# ----------------------------------------------------------------------------
# IW28 — notas de manutenção
# ----------------------------------------------------------------------------
_NOTAS_COLUNAS = {
    "Nota": [r"^NOTA$"], "Ordem": [r"^ORDEM$"], "Tipo de nota": [r"TIPO DE NOTA"],
    "Descrição": [r"^DESCRICAO$", r"TEXTO BREVE", r"DESCRICAO"], "Equipamento": [r"^EQUIPAMENTO$"],
    "Objeto técnico": [r"DENOMINACAO DO OBJETO", r"OBJETO TECNICO"], "Local de instalação": [r"DENOMINACAO DO LOC"],
    "Centro de trabalho": [r"CENTRO TRAB RESPONS", r"^CENTRO DE TRABALHO$"], "Localização": [r"^LOCALIZACAO$"],
    "Notificador": [r"NOTIFICADOR"], "Parada": [r"^PARADA$"], "Duração da parada": [r"DURACAO DA PARADA"],
    "Início avaria": [r"^INICIO (DA )?AVARIA$", r"^(DATA )?INICIO (DA )?AVARIA"],
    "Fim avaria": [r"^FIM (DA )?AVARIA$", r"^(DATA )?FIM (DA )?AVARIA"],
    "Data da nota": [r"DATA DA NOTA"], "Criado por": [r"CRIADO POR"], "Código ABC": [r"CODIGO ABC", r"^ABC$"],
    "Plano": [r"PLANO DE MANUTENCAO"],
}


def preparar_notas(cru: pd.DataFrame, hoje: pd.Timestamp | None = None, excluir_piloto_matriz: bool = True) -> pd.DataFrame:
    hoje = (hoje or agora_local()).normalize()
    df = _padronizar(cru, _NOTAS_COLUNAS, ["Nota", "Ordem", "Tipo de nota", "Descrição", "Equipamento", "Objeto técnico",
                                            "Local de instalação", "Centro de trabalho", "Localização", "Notificador",
                                            "Parada", "Criado por", "Código ABC", "Plano"])
    df = df[df["Nota"] != ""].copy()
    if excluir_piloto_matriz:
        df = df[~fora_da_visao(df)]
    for c in ["Início avaria", "Fim avaria", "Data da nota"]:
        df[c] = para_data(df[c])
    df["Com parada"] = df["Parada"].str.upper().isin(["X", "SIM", "S", "1"])
    df["Horas parado"] = para_numero(df["Duração da parada"]).fillna(0.0)
    df["Com ordem"] = df["Ordem"] != ""
    df["Data"] = df["Data da nota"].fillna(df["Início avaria"])
    df["Dias"] = (hoje - df["Data"]).dt.days
    df["Equip. (chave)"] = np.where(df["Equipamento"] != "", df["Equipamento"], df["Objeto técnico"])
    df["Setor"] = setores(df["Centro de trabalho"])
    return df.drop(columns=["Parada", "Duração da parada"]).reset_index(drop=True)


@st.cache_resource(show_spinner="Preparando notas (IW28)…", max_entries=2)
def _notas(origens: tuple, excluir: bool, dia: str) -> pd.DataFrame:
    return preparar_notas(_cru(NOTAS, origens), pd.Timestamp(dia), excluir_piloto_matriz=excluir)


# ----------------------------------------------------------------------------
# IH08 — cadastro de equipamentos
# ----------------------------------------------------------------------------
ABC_PARA_CRITICIDADE = {"A": "Alta", "B": "Média", "C": "Baixa"}


def preparar_equipamentos(cru: pd.DataFrame, excluir_piloto_matriz: bool = True) -> pd.DataFrame:
    df = _padronizar(cru, {"Equipamento": [r"^EQUIPAMENTO$"], "Denominação": [r"DENOMINACAO"],
                           "Localização": [r"^LOCALIZACAO$"], "Código ABC": [r"CODIGO ABC", r"^ABC$"],
                           "Local de instalação": [r"LOCAL DE INSTALACAO", r"LOC INSTAL"]},
                     ["Equipamento", "Denominação", "Localização", "Código ABC", "Local de instalação"])
    df = df[df["Equipamento"] != ""].drop_duplicates("Equipamento").copy()
    if excluir_piloto_matriz:
        df = df[~fora_da_visao(df)].copy()
    df["Código ABC"] = df["Código ABC"].str.upper().str.strip()
    df["Desativado"] = df["Denominação"].map(chave).str.contains(r"DESATIV|INATIV|NAO ESTA NA FABRICA|SUCATA|BAIXAD")
    return df.reset_index(drop=True)


@st.cache_resource(show_spinner="Preparando cadastro de equipamentos (IH08)…", max_entries=2)
def _equip(origens: tuple, excluir: bool) -> pd.DataFrame:
    return preparar_equipamentos(_cru(EQUIP, origens), excluir)


# ----------------------------------------------------------------------------
# IW47 — confirmações (horas apontadas por pessoa)
# ----------------------------------------------------------------------------
_CONF_COLUNAS = {
    "Nº pessoal": [r"^N PESSOAL$", r"NUMERO PESSOAL"], "Nome": [r"NOME DO EMPREGADO", r"^NOME$"],
    "Ordem": [r"^ORDEM$"], "Operação": [r"^OPERACAO$"], "Status sistema": [r"STATUS DO SISTEMA"],
    "Trabalho": [r"^TRABALHO REAL$"], "Unidade": [r"UNID TRABALHO REAL", r"UNID.*TRAB"],
    "Centro de trabalho": [r"CENTRO TRAB REAL", r"^CENTRO DE TRABALHO$", r"CENTRO TRAB"],
    "Atividade": [r"TP ATIVIDADE REAL", r"TIPO ATIVID.*REAL"], "Atividade planejada": [r"TIPO ATIVID.*PLAN"],
    "Data": [r"DATA LANCAMENTO", r"DATA DO FIM REAL", r"CRIADO EM"],
    "Tipo": [r"^TIPO ORDEM$", r"^TIPO DE ORDEM$"], "Equipamento": [r"^EQUIPAMENTO$"],
    "Texto": [r"TEXTO DE CONFIRMACAO"], "Causa do desvio": [r"CAUSA DO DESVIO"],
}


def preparar_confirmacoes(cru: pd.DataFrame, excluir_piloto_matriz: bool = True) -> pd.DataFrame:
    df = _padronizar(cru, _CONF_COLUNAS, ["Nº pessoal", "Nome", "Ordem", "Operação", "Status sistema", "Unidade",
                                           "Centro de trabalho", "Atividade", "Atividade planejada", "Tipo",
                                           "Equipamento", "Texto", "Causa do desvio"])
    df = df[(df["Ordem"] != "") & (df["Nº pessoal"] != "")].copy()
    if excluir_piloto_matriz:
        df = df[~fora_da_visao(df)]
    fator = df["Unidade"].str.upper().map(_HORAS_POR).fillna(1 / 60)
    # estornos vêm com trabalho negativo e anulam o apontamento original: a soma fica certa
    df["Horas"] = (para_numero(df["Trabalho"]).fillna(0.0) * fator).round(3)
    df["Data"] = para_data(df["Data"])
    df["Atividade"] = df["Atividade"].where(df["Atividade"] != "", df["Atividade planejada"])
    df["Setor"] = setores(df["Centro de trabalho"])
    return df.drop(columns=["Trabalho", "Atividade planejada"]).reset_index(drop=True)


def so_da_equipe(conf: pd.DataFrame | None, equipe: pd.DataFrame | None) -> pd.DataFrame | None:
    """Apontamentos (IW47) só das pessoas da planilha de Gestão de HH — a base das visões de HH (horas,
    utilização, quem fez o quê). Sem a planilha, não há como filtrar: volta tudo."""
    if conf is None or equipe is None or not len(equipe) or "Nº pessoal" not in conf:
        return conf
    return conf[conf["Nº pessoal"].isin(set(equipe["Nº pessoal"]))]


@st.cache_resource(show_spinner="Preparando apontamentos de horas (IW47)…", max_entries=2)
def _conf(origens: tuple, excluir: bool) -> pd.DataFrame:
    return preparar_confirmacoes(_cru(CONF, origens), excluir)


# ----------------------------------------------------------------------------
# Gestão de HH — dados das pessoas (só a lista da equipe)
# ----------------------------------------------------------------------------
_EQUIPE_COLUNAS = {
    "Nº pessoal": [r"^N PESSOAL$", r"NUMERO PESSOAL", r"MATRICULA"], "Nome": [r"^NOME$", r"^NOME"],
    "Cargo": [r"^CARGO$", r"^FUNCAO$"], "Centro de trabalho": [r"^CENTRO TRABALHO$", r"CENTRO TRAB"],
    "Área": [r"^AREA$"], "Turma": [r"^TURMA$", r"^TURNO$"], "Supervisor": [r"SUPERVISOR"],
}


def preparar_equipe(cru: pd.DataFrame, excluir_piloto_matriz: bool = True) -> pd.DataFrame:
    df = _padronizar(cru, _EQUIPE_COLUNAS, list(_EQUIPE_COLUNAS))
    if excluir_piloto_matriz:
        df = df[~fora_da_visao(df)]
    df = df[df["Nº pessoal"].str.fullmatch(r"\d{3,}") & (df["Nome"] != "")]
    df = df[~df["Nome"].str.upper().str.startswith("#")]  # fórmulas quebradas (#REF!, #N/A)
    for c in ["Cargo", "Área", "Turma", "Supervisor"]:
        df[c] = df[c].where(~df[c].str.startswith("#"), "")
    df["Especialidade"] = df["Cargo"].map(_especialidade)
    df["Setor"] = setores(df["Centro de trabalho"])
    return df.drop_duplicates("Nº pessoal", keep="last").reset_index(drop=True)


def _especialidade(cargo: str) -> str:
    k = chave(cargo)
    for rx, nome in [(r"MATRIZ", "Matrizaria"), (r"^GPM$|PLANEJ|PROGRAMAD", "Planejamento (GPM)"),
                     (r"ELETROMEC", "Eletromecânica"), (r"ELETRIC|ELETRON", "Elétrica"), (r"MECANIC", "Mecânica"),
                     (r"LUBRIF", "Lubrificação"), (r"INSTRUM|AUTOMA", "Instrumentação"), (r"SOLDA|CALDEIR", "Caldeiraria"),
                     (r"ANALISTA|ASSIST|SUPERV|COORD|ENGENH|SPV|ANL|AST", "Apoio / gestão")]:
        if re.search(rx, k):
            return nome
    return "Outros"


@st.cache_resource(show_spinner="Preparando equipe (Gestão de HH)…", max_entries=2)
def _equipe(origens: tuple, excluir: bool) -> pd.DataFrame:
    return preparar_equipe(_cru(EQUIPE, origens), excluir)


@st.cache_data(show_spinner=False, max_entries=4)
def _tipos_da_gestao(id: str, assinatura: str) -> dict[str, str]:
    """Nomes dos tipos de ordem (YM11 Corretiva Emergencial…) da aba de apoio da planilha de Gestão de HH."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(_bytes(id, assinatura)), read_only=True, data_only=True)
        for ws in wb.worksheets:
            linhas = list(ws.iter_rows(max_row=60, values_only=True))
            if not linhas:
                continue
            cab = [chave(c) for c in linhas[0]]
            if "TIPO DE OM" in cab and "TIPO" in cab:
                i, j = cab.index("TIPO"), cab.index("TIPO DE OM")
                out = {str(r[i]).strip(): str(r[j]).strip() for r in linhas[1:]
                       if len(r) > max(i, j) and r[i] and r[j] and re.fullmatch(r"[A-Z]{1,3}\d{2}", str(r[i]).strip())}
                wb.close()
                return out
        wb.close()
    except Exception:  # noqa: BLE001 — planilha sem a aba de apoio
        pass
    return {}


def tipos_de_ordem_padrao() -> dict[str, str]:
    b = equipe()
    if b.origem is None or not b.origem.arquivo.arquivo.lower().endswith((".xlsx", ".xlsm")):
        return {}
    return _tipos_da_gestao(b.origem.arquivo.id, b.origem.arquivo.assinatura)


# ----------------------------------------------------------------------------
# Gerenciador de AF — análises de falha e ações (ver central/af.py)
# ----------------------------------------------------------------------------
@st.cache_resource(show_spinner="Preparando análises de falha (Gerenciador de AF)…", max_entries=2)
def _afs(origens: tuple, dia: str) -> pd.DataFrame:
    return af_mod.preparar_afs(_cru(AF, origens), pd.Timestamp(dia))


@st.cache_resource(show_spinner="Preparando ações das análises de falha…", max_entries=2)
def _acoes_af(origens: tuple, origens_af: tuple, dia: str) -> pd.DataFrame:
    afs_df = _afs(origens_af, dia) if origens_af else None
    return af_mod.preparar_acoes(_cru(ACAO_AF, origens), pd.Timestamp(dia), afs_df)


# ----------------------------------------------------------------------------
# Cadastros feitos no próprio app (JSON na pasta do app, compartilhada)
# ----------------------------------------------------------------------------
@st.cache_data(show_spinner=False, max_entries=16)
def _ler_json(id: str, assinatura: str) -> dict:
    """Erro de leitura levanta (e não fica em cache): quem chama usa a última leitura boa."""
    dados = json.loads(fonte().ler(id).decode("utf-8"))
    return dados if isinstance(dados, dict) else {}


_ULTIMO_BOM: dict[str, dict] = {}   # id do cadastro → última leitura boa (se a fonte falhar um instante)


def _ler_json_seguro(id: str, assinatura: str) -> dict:
    try:
        dados = _ler_json(id, assinatura)
    except Exception:  # noqa: BLE001 — rede/arquivo bloqueado: não some com o cadastro (nem desloga ninguém)
        return _ULTIMO_BOM.get(id, {})
    _ULTIMO_BOM[id] = dados
    return dados


def _ler_json_agora(nome: str) -> dict:
    """Relê o cadastro agora. Só devolve {} quando o arquivo não existe: uma falha de leitura
    (rede, GitHub/SharePoint fora do ar) levanta erro — tratá-la como "vazio" e gravar em cima
    apagaria o cadastro inteiro (foi assim que a lista de usuários se perdeu)."""
    try:
        bruto = fonte().ler(fonte().id_de(nome))
    except fontes.NaoEncontrado:
        return {}
    except Exception as e:  # noqa: BLE001
        raise fontes.FonteErro(f"Não foi possível ler {nome} agora ({e}). Nada foi gravado; "
                               "tente de novo em instantes.") from e
    try:
        dados = json.loads(bruto.decode("utf-8"))
    except ValueError as e:
        raise fontes.FonteErro(f"{nome} está ilegível; nada foi gravado para não perder o conteúdo.") from e
    if not isinstance(dados, dict):
        raise fontes.FonteErro(f"{nome} está em formato inesperado; nada foi gravado.")
    return dados


@st.cache_resource
def _backups_feitos() -> set:
    return set()


def _backup_diario(nome: str, conteudo: dict) -> None:
    """Uma cópia por dia de cada cadastro (antes da primeira gravação do dia) em backup/."""
    dia = hoje_local().isoformat()
    if not conteudo or (nome, dia) in _backups_feitos():
        return
    try:
        base = nome.rsplit(".", 1)[0].replace("/", "_")
        fonte().gravar(f"backup/{base}-{dia}.json", json.dumps(conteudo, ensure_ascii=False, indent=1).encode("utf-8"))
        _backups_feitos().add((nome, dia))
    except Exception:  # noqa: BLE001 — a cópia de segurança nunca impede a gravação
        pass


def manifesto() -> dict:
    """Manifesto do repositório de dados (modo GitHub): de onde e quando veio cada base."""
    arq = next((a for a in arquivos() if a.id == fontes.MANIFESTO), None)
    return _ler_json_seguro(arq.id, arq.assinatura) if arq else {}


def ler_cadastro(nome: str) -> dict:
    alvo = fonte().id_de(nome)
    arq = next((a for a in arquivos() if a.id == alvo), None)
    return _ler_json_seguro(arq.id, arq.assinatura) if arq else {}


def gravar_cadastro(nome: str, chave_item: str, dados: dict | None, usuario: str = "") -> None:
    gravar_cadastro_lote(nome, {chave_item: dados}, usuario)


_CAMPOS_DE_EDICAO = ("atualizado_em", "atualizado_por")


_TRAVAS: dict[str, threading.Lock] = {}


def _trava(nome: str) -> threading.Lock:
    return _TRAVAS.setdefault(nome, threading.Lock())


def _decodificar_cadastro(nome: str, bruto: bytes | None) -> dict:
    if bruto is None:
        return {}
    try:
        dados = json.loads(bruto.decode("utf-8"))
    except ValueError as e:
        raise fontes.FonteErro(f"{nome} está ilegível; nada foi gravado para não perder o conteúdo.") from e
    if not isinstance(dados, dict):
        raise fontes.FonteErro(f"{nome} está em formato inesperado; nada foi gravado.")
    return dados


def _ler_versionado(nome: str) -> tuple[bytes | None, str | None]:
    f = fonte()
    if hasattr(f, "ler_versionado"):
        return f.ler_versionado(nome)
    try:                                   # fonte sem versão: lê e grava por cima (como antes)
        return f.ler(f.id_de(nome)), None
    except fontes.NaoEncontrado:
        return None, None


def _gravar_condicional(nome: str, conteudo: bytes, versao: str | None) -> bool:
    f = fonte()
    if hasattr(f, "gravar_condicional"):
        return f.gravar_condicional(nome, conteudo, versao)
    f.gravar(nome, conteudo)
    return True


def gravar_cadastro_lote(nome: str, itens: dict[str, dict | None], usuario: str = "", mesclar: bool = False) -> None:
    """Atualiza só os itens informados (relendo o arquivo na hora), para duas pessoas
    editando itens diferentes não se sobrescreverem. Valor None remove o item.

    A gravação é condicional à versão lida (sha no GitHub, eTag no SharePoint, data no disco): se outra
    pessoa (ou o robô) gravou no meio, relê e aplica de novo — nenhuma das duas alterações se perde.

    mesclar=True muda só os campos informados de cada item (campo None é apagado) e
    remove o item que fica sem nenhum campo — ex.: o mínimo de um material não apaga a foto."""
    with _trava(nome):
        for tentativa in range(5):
            try:
                bruto, versao = _ler_versionado(nome)
            except fontes.FonteErro as e:
                raise fontes.FonteErro(f"Não foi possível ler {nome} agora ({e}). Nada foi gravado; "
                                       "tente de novo em instantes.") from e
            atual = _decodificar_cadastro(nome, bruto)   # erro de leitura nunca vira "vazio" gravado por cima
            if tentativa == 0:
                _backup_diario(nome, atual)
            agora = datetime.now(timezone.utc).isoformat(timespec="seconds")
            for chave_item, dados in itens.items():
                if dados is not None and mesclar:
                    dados = {k: v for k, v in {**(atual.get(chave_item) or {}), **dados}.items()
                             if v is not None and k not in _CAMPOS_DE_EDICAO}
                    dados = dados or None
                if dados is None:
                    atual.pop(chave_item, None)
                else:
                    atual[chave_item] = {**dados, "atualizado_em": agora, "atualizado_por": usuario}
            conteudo = json.dumps(atual, ensure_ascii=False, indent=1).encode("utf-8")
            if _gravar_condicional(nome, conteudo, versao):
                break
            time.sleep(0.3 * (tentativa + 1))
        else:
            raise fontes.FonteErro(f"{nome} está sendo alterado por outra pessoa agora; nada foi gravado — "
                                   "tente de novo em instantes.")
    recarregar()


# ----------------------------------------------------------------------------
# Fotos dos materiais (ver central/fotos.py)
# ----------------------------------------------------------------------------
@st.cache_data(show_spinner=False, max_entries=128)
def _foto(caminho: str, versao: str) -> bytes | None:
    try:
        return fonte().ler(fonte().id_de(caminho))
    except Exception:  # noqa: BLE001 — foto apagada da pasta à mão
        return None


def foto_material(info: dict | None) -> bytes | None:
    """Foto do material a partir do item do cadastro de materiais (None se não tem)."""
    if not info or not info.get("foto"):
        return None
    return _foto(info["foto"], info.get("foto_versao", ""))


@st.cache_data(show_spinner=False, max_entries=1024)
def miniatura_material(caminho: str, versao: str) -> str | None:
    conteudo = _foto(caminho, versao)
    try:
        return fotos.miniatura(conteudo) if conteudo else None
    except fotos.FotoErro:
        return None


def gravar_foto_material(codigo: str, conteudo: bytes, usuario: str = "") -> None:
    jpeg = fotos.preparar(conteudo)
    caminho = fotos.caminho(codigo)
    fonte().gravar(caminho, jpeg)
    gravar_cadastro_lote(ARQ_CAD_MAT, {codigo: {"foto": caminho, "foto_versao": fotos.versao(jpeg)}}, usuario,
                         mesclar=True)


def remover_foto_material(codigo: str, usuario: str = "") -> None:
    """Desvincula a foto (o arquivo fica na pasta; uma foto nova o substitui)."""
    gravar_cadastro_lote(ARQ_CAD_MAT, {codigo: {"foto": None, "foto_versao": None}}, usuario, mesclar=True)


def enviar_arquivo(tipo: str, nome_original: str, conteudo: bytes) -> str:
    """Valida e grava na pasta do app um arquivo exportado, já com o tipo no nome."""
    tipos = {t for t, _, _ in leitura.sondar(nome_original, conteudo)}
    if tipo not in tipos:
        esperado = leitura.NOMES_BASE[tipo]
        outro = next(iter(tipos), None)
        raise leitura.LeituraErro(
            f"o arquivo não parece ser de {esperado}"
            + (f" (parece {leitura.NOMES_BASE[outro]})" if outro else " — confira as colunas do export"))
    ext = PurePath(nome_original).suffix.lower() or ".xlsx"
    prefixo = {IW38: "IW38", MB52: "MB52", REQ: "REQUISICOES", IP19: "IP19", OPER: "IW38OP", NOTAS: "IW28",
               EQUIP: "IH08", CONF: "IW47", EQUIPE: "GESTAO_HH", AF: "GERENCIADOR_AF",
               ACAO_AF: "GERENCIADOR_AF_ACOES"}[tipo]
    nome = f"{prefixo}_{datetime.now():%Y-%m-%d_%H%M%S}{ext}"
    fonte().gravar(nome, conteudo)
    recarregar()
    return nome


def recarregar() -> None:
    _listar.clear()
    # o arquivo mudou por gravação feita nesta sessão: o vigia da barra lateral não deve tratar isso
    # como "arquivo novo" e reiniciar o app (o que apagaria o estado dos widgets da página)
    try:
        st.session_state["_assinatura"] = None
    except Exception:  # noqa: BLE001 — fora de uma sessão (testes, sincronizador)
        pass


# ----------------------------------------------------------------------------
# Acesso de alto nível usado pelas páginas
# ----------------------------------------------------------------------------
@dataclass
class Base:
    df: pd.DataFrame | None
    origem: Origem | None          # arquivo principal (o mais novo)
    erro: str = ""
    extra: object = None           # IW38: nº removidas · MB52: detalhe por depósito
    origens: tuple = ()            # todos os arquivos somados

    @property
    def atualizado(self) -> datetime | None:
        return self.origem.arquivo.modificado if self.origem else None


def _carregar(tipo: str, fn) -> Base:
    usados = tuple(inventario().usados.get(tipo, []))
    if not usados:
        return Base(None, None)
    try:
        b = fn(usados)
        b.origem, b.origens = usados[0], usados
        return b
    except Exception as e:  # noqa: BLE001
        return Base(None, usados[0], erro=str(e), origens=usados)


def _excluir() -> bool:
    return config.carregar().excluir_piloto_matriz


def iw38() -> Base:
    def fn(us):
        df, removidas = _iw38(us, _excluir(), str(hoje_local()))
        return Base(df, None, extra=removidas)
    return _carregar(IW38, fn)


def mb52() -> Base:
    def fn(us):
        resumo, det = _mb52(us)
        return Base(resumo, None, extra=det)
    return _carregar(MB52, fn)


def requisicoes() -> Base:
    return _carregar(REQ, lambda us: Base(_req(us, str(hoje_local())), None))


def afs() -> Base:
    """Análises de falha (a situação depende do dia: o cache vira à meia-noite)."""
    return _carregar(AF, lambda us: Base(_afs(us, str(hoje_local())), None))


def acoes_af() -> Base:
    origens_af = tuple(inventario().usados.get(AF, []))
    return _carregar(ACAO_AF, lambda us: Base(_acoes_af(us, origens_af, str(hoje_local())), None))


def confirmacoes() -> Base:
    return _carregar(CONF, lambda us: Base(_conf(us, _excluir()), None))


def equipe() -> Base:
    return _carregar(EQUIPE, lambda us: Base(_equipe(us, _excluir()), None))


def ip19() -> Base:
    return _carregar(IP19, lambda us: Base(_ip19(us, _excluir()), None))


def operacoes() -> Base:
    return _carregar(OPER, lambda us: Base(_oper(us, _excluir()), None))


def tipos_por_ordem() -> pd.Series:
    """Ordem → tipo de ordem (YM11, YM13…) juntando IW38, IW38OP e IW47."""
    partes = []
    for b in (iw38(), operacoes(), confirmacoes()):
        df = b.df
        if df is not None and len(df) and "Tipo" in df:
            t = pd.DataFrame({"Ordem": df["Ordem"].astype(str), "Tipo": df["Tipo"].astype(object).fillna("").astype(str)})
            partes.append(t[t["Tipo"].str.strip() != ""])
    if not partes:
        return pd.Series(dtype=str)
    t = pd.concat(partes, ignore_index=True).drop_duplicates("Ordem")
    return t.set_index("Ordem")["Tipo"].str.strip().str.upper()


@st.cache_resource(show_spinner="Marcando as quebras (Y1 → YM11)…", max_entries=2)
def _notas_quebras(origens: tuple, chave_tipos: tuple, dia: str) -> pd.DataFrame:
    return marcar_quebras(_notas(origens, _excluir(), dia), tipos_por_ordem())


def notas() -> Base:
    """IW28 com a coluna Quebra (nota Y1 convertida em ordem YM11)."""
    usados = inventario().usados
    chave_tipos = tuple(tuple(usados.get(t, [])) for t in (IW38, OPER, CONF)) + (_excluir(),)
    return _carregar(NOTAS, lambda us: Base(_notas_quebras(us, chave_tipos, str(hoje_local())), None))


def equipamentos() -> Base:
    return _carregar(EQUIP, lambda us: Base(_equip(us, _excluir()), None))

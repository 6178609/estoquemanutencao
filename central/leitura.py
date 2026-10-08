"""Leitura dos arquivos exportados do SAP e de planilhas, em qualquer formato
usual: .xlsx, .xls (inclusive o "xls" que na verdade é texto/HTML do SAP GUI),
.csv, .txt (lista com |), .htm/.html ("salvar como HTML" do SAP) e .parquet.

Também reconhece sozinho de qual base é o arquivo (IW38, MB52 ou requisições)
pelos cabeçalhos — não depende do nome do arquivo.
"""

from __future__ import annotations

import io
import re
from pathlib import PurePath

import pandas as pd

from .util import chave

EXTENSOES = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".htm", ".html", ".parquet")

IW38, MB52, REQ, IP19 = "iw38", "mb52", "requisicoes", "ip19"
OPER, NOTAS, EQUIP = "iw38op", "iw28", "ih08"
CONF, EQUIPE = "iw47", "equipe"
AF, ACAO_AF = "af", "acoes_af"
NOMES_BASE = {IW38: "Ordens (IW38)", IP19: "Planos de manutenção (IP19)", OPER: "Operações das ordens (IW38OP)",
              NOTAS: "Notas de manutenção (IW28)", EQUIP: "Cadastro de equipamentos (IH08)",
              MB52: "Estoque (MB52)", REQ: "Requisições de compra",
              CONF: "Apontamentos de horas (IW47)", EQUIPE: "Equipe de manutenção (Gestão de HH)",
              AF: "Análises de falha (Gerenciador de AF)", ACAO_AF: "Ações das análises de falha (Gerenciador de AF)"}
TIPOS = (IW38, IP19, OPER, NOTAS, EQUIP, MB52, REQ, CONF, EQUIPE, AF, ACAO_AF)

# Cada base é UM arquivo: vale o export mais recente (o PCM gera um IW38 único com plano
# de manutenção e backlog, e um IW28 único). O mecanismo abaixo permite que uma base
# some vários arquivos pela chave (vale o mais novo para a mesma chave; entram os
# MAX_MESCLA mais recentes e os com BK/HIST no nome) — hoje nenhuma base usa.
MESCLAR: dict[str, str] = {}
MAX_MESCLA = 6
_HISTORICO = re.compile(r"BK|HIST", re.I)


class LeituraErro(ValueError):
    pass


# ----------------------------------------------------------------------------
# Texto
# ----------------------------------------------------------------------------
def _decodificar(conteudo: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return conteudo.decode(enc)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1", errors="replace")


def _ler_texto(txt: str) -> pd.DataFrame:
    linhas = [l for l in txt.splitlines() if l.strip()]
    if not linhas:
        raise LeituraErro("arquivo vazio")
    amostra = "\n".join(linhas[:50])

    # Lista do SAP "não convertida": | col | col | e linhas de ------
    if amostra.count("|") > len(linhas[:50]) * 2:
        uteis = [l for l in linhas if "|" in l and not re.fullmatch(r"[\s|\-]+", l)]
        tabela = [[c.strip() for c in l.strip().strip("|").split("|")] for l in uteis]
        return _tabela_com_cabecalho(tabela)

    sep = max([";", "\t", ","], key=amostra.count)
    df = pd.read_csv(io.StringIO("\n".join(linhas)), sep=sep, dtype=str, keep_default_na=False)
    return df


def _tabela_com_cabecalho(tabela: list[list[str]]) -> pd.DataFrame:
    """Primeira linha com 3+ células preenchidas vira cabeçalho; repetições do cabeçalho
    (o SAP repete a cada página) são descartadas."""
    idx = next((i for i, l in enumerate(tabela) if sum(1 for c in l if c) >= 3), None)
    if idx is None:
        raise LeituraErro("não encontrei uma linha de cabeçalho")
    cab = [c or f"col{i}" for i, c in enumerate(tabela[idx])]
    # nomes repetidos ganham sufixo
    vistos: dict[str, int] = {}
    for i, c in enumerate(cab):
        if c in vistos:
            vistos[c] += 1
            cab[i] = f"{c} ({vistos[c]})"
        else:
            vistos[c] = 0
    linhas = []
    for l in tabela[idx + 1:]:
        if l[: len(cab)] == tabela[idx][: len(cab)] or not any(l):
            continue
        l = (l + [""] * len(cab))[: len(cab)]
        linhas.append(l)
    return pd.DataFrame(linhas, columns=cab)


# ----------------------------------------------------------------------------
# HTML do SAP GUI
# ----------------------------------------------------------------------------
def _ler_html(txt: str) -> pd.DataFrame:
    from lxml import html as lhtml

    doc = lhtml.fromstring(txt)
    tabelas = doc.xpath("//table")
    if not tabelas:
        raise LeituraErro("nenhuma tabela no HTML")
    principal = max(tabelas, key=lambda t: len(t.xpath(".//tr")))
    linhas = [
        [re.sub(r"\s+", " ", (td.text_content() or "").replace("\xa0", " ")).strip() for td in tr.xpath("./th|./td")]
        for tr in principal.xpath(".//tr")
    ]
    largura = max((len(l) for l in linhas), default=0)
    if largura >= 3:
        return _tabela_com_cabecalho(linhas)
    # Lista agrupada (MB52 em árvore): cada linha é um texto de largura fixa.
    brutas = [tr.text_content().replace("\xa0", " ") for tr in principal.xpath(".//tr")]
    return _lista_agrupada_mb52(brutas)


def _tokens(linha: str) -> list[str]:
    return [t for t in re.split(r"\s{2,}", linha.strip()) if t]


def _num_br(t: str):
    """Número no padrão do SAP (1.234,500); o sinal pode vir no fim (12,500- = −12,5)."""
    t = t.strip()
    negativo = t.endswith("-")
    try:
        v = float(t.rstrip("-").replace(".", "").replace(",", "."))
    except ValueError:
        return None
    return -v if negativo else v


_QTD = re.compile(r"-?\d[\d.]*(,\d+)?-?")
_UNIDADE = re.compile(r"[A-Za-zÀ-ÿ]{1,3}\d?")
_TOTAL = re.compile(r"^\s*\*+\s*total", re.IGNORECASE)


def _lista_agrupada_mb52(linhas: list[str]) -> pd.DataFrame:
    """Lista do MB52 salva como HTML pelo SAP GUI, em árvore:

        48737      LAMPADA V MET OVOIDE LEIT E40 400W        ← material (na coluna 0: código + texto)
                   UN                                  6     ← saldo (recuado: UM + quantidade)
        84518      PROD QUIM PRT 601 07
                   KG                          2.475,000
        0099378    PAR   A                         3.442     ← saldo de um lote (coluna 0: nº + UM + quantidade)
        * Total                                              ← total geral (ignorado)

    A linha é classificada pela posição, não pelo "jeito de número": o texto do material pode ter números
    (CONJUNTO VEDACAO  2690089  2660019) e não pode virar saldo do material anterior. Os saldos de cada
    material são somados."""
    grupos, atual = [], None
    for linha in linhas:
        toks = _tokens(linha)
        if not toks or _TOTAL.match(linha):
            continue
        recuada = linha[:1].isspace()
        # lote/avaliação na coluna 0: nº, UM e uma quantidade depois (o depósito pode vir no fim)
        lote = (not recuada and len(toks) >= 3 and _UNIDADE.fullmatch(toks[1]) is not None
                and any(_QTD.fullmatch(t) for t in toks[2:]))
        if not recuada and re.match(r"^\d", toks[0]) and not lote:
            atual = {"Material": toks[0], "Texto breve material": " ".join(toks[1:]), "UM": "", "Utilização livre": 0.0}
            grupos.append(atual)
            continue
        qtds = [t for t in toks[1:] if _QTD.fullmatch(t)]
        if atual is None or not qtds:
            continue
        q = _num_br(qtds[-1])
        if q is None:
            continue
        atual["Utilização livre"] += q
        if not atual["UM"]:
            atual["UM"] = next((t for t in toks if _UNIDADE.fullmatch(t) and not t.isdigit()), "")
    if not grupos:
        raise LeituraErro("lista do SAP sem materiais reconhecíveis")
    return pd.DataFrame(grupos)


# ----------------------------------------------------------------------------
# Entrada principal
# ----------------------------------------------------------------------------
def _eh_xlsx(conteudo: bytes) -> bool:
    return conteudo[:4] == b"PK\x03\x04"


def _eh_xls(conteudo: bytes) -> bool:
    return conteudo[:4] == b"\xd0\xcf\x11\xe0"


def ler_arquivo(nome: str, conteudo: bytes, planilha: str | None = None, linha_cab: int | None = None) -> pd.DataFrame:
    """Lê o arquivo inteiro. Para Excel, `planilha` e `linha_cab` (0 = primeira linha)
    vêm de `sondar`; sem eles usa a primeira aba e a primeira linha."""
    ext = PurePath(nome).suffix.lower()
    inicio = conteudo[:2048].lstrip().lower()
    try:
        if ext == ".parquet":
            df = pd.read_parquet(io.BytesIO(conteudo))
        elif _eh_xlsx(conteudo) or _eh_xls(conteudo):
            df = pd.read_excel(io.BytesIO(conteudo), engine="openpyxl" if _eh_xlsx(conteudo) else "xlrd",
                               sheet_name=planilha or 0, header=linha_cab or 0, dtype=object)
        elif b"<table" in conteudo[:200_000].lower() or inicio.startswith((b"<!doctype", b"<html")):
            df = _ler_html(_decodificar(conteudo))
        else:
            df = _ler_texto(_decodificar(conteudo))
    except LeituraErro:
        raise
    except Exception as e:  # noqa: BLE001 — qualquer falha de parser vira mensagem amigável
        raise LeituraErro(f"não foi possível ler {nome}: {e}") from e

    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[:, [c for c in df.columns if c and not c.lower().startswith("unnamed")]]
    df = df.dropna(how="all")
    if df.empty:
        raise LeituraErro(f"{nome} não tem linhas de dados")
    return df.reset_index(drop=True)


def sondar(nome: str, conteudo: bytes) -> list[tuple[str, str | None, int | None]]:
    """Descobre quais bases existem no arquivo, sem ler tudo quando é Excel:
    devolve [(tipo, aba, linha do cabeçalho)]. Uma pasta de trabalho pode ter
    mais de uma base (ex.: uma aba IW38 e outra MB52)."""
    if eh_recorte(nome):
        return []
    achados = []
    if _eh_xlsx(conteudo) or _eh_xls(conteudo):
        try:
            if _eh_xlsx(conteudo):
                from openpyxl import load_workbook

                wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
                abas = [(ws.title, list(ws.iter_rows(max_row=15, values_only=True)),
                         getattr(ws, "sheet_state", "visible") == "visible") for ws in wb.worksheets]
                wb.close()
            else:
                folhas = pd.read_excel(io.BytesIO(conteudo), engine="xlrd", sheet_name=None, header=None, nrows=15)
                abas = [(t, f.values.tolist(), True) for t, f in folhas.items()]
        except Exception:  # noqa: BLE001
            return []
        visiveis = []
        for titulo, linhas, visivel in abas:
            for i, linha in enumerate(linhas):
                tipo = identificar([str(c) for c in linha if c is not None and str(c).strip()])
                if tipo:
                    achados.append((tipo, titulo, i))
                    visiveis.append(visivel)
                    break
        # uma aba por base em cada arquivo: a visível vence a oculta (ex.: "Análise de Falha" e a
        # "Análise de Falha (Backup)" escondida no Gerenciador de AF)
        ordem = sorted(range(len(achados)), key=lambda k: not visiveis[k])
        vistos, unicos = set(), []
        for k in ordem:
            if achados[k][0] not in vistos:
                vistos.add(achados[k][0])
                unicos.append(achados[k])
        achados = unicos
        # Planilhas de gestão trazem cópias antigas de outras bases: da Gestão de HH só vale a lista de
        # pessoas (aba IW47, "DadosExtras (IW38)"… ficam de fora) e do Gerenciador de AF só as análises e
        # as ações (a aba "Gargalos" parece um IH08) — as demais bases vêm dos exports do SAP.
        for proprios in ({EQUIPE}, {AF, ACAO_AF}):
            if any(t in proprios for t, _, _ in achados):
                achados = [a for a in achados if a[0] in proprios]
        return achados
    try:
        tipo = identificar(ler_arquivo(nome, conteudo).columns)
    except LeituraErro:
        return []
    return [(tipo, None, None)] if tipo else []


# Colunas que só existem nas planilhas baixadas do próprio site (botão "Baixar Excel").
# Elas parecem um export do SAP, mas são recortes filtrados: nunca podem virar a base.
_COLUNAS_DO_APP = r"^(NATUREZA|SITUACAO|DIAS EM ABERTO|DIAS AGUARDANDO|PECAS EM FALTA|USADO EM|COM PLANO|ATRASADA)$"


def eh_recorte(nome: str) -> bool:
    """Exports filtrados (deste site ou do site antigo: ordens_filtradas, IW38_filtrado.csv…)."""
    return bool(re.search(r"filtrad", PurePath(nome).stem, re.I))


def escolher_origens(tipo: str, candidatos: list) -> list:
    """Dos candidatos de uma base (do mais novo para o mais velho; cada um com
    `.arquivo.nome`), quais usar: o mais novo, ou — nas bases que se somam — os
    MAX_MESCLA mais novos + os arquivos de histórico."""
    if not candidatos:
        return []
    if tipo not in MESCLAR:
        return candidatos[:1]
    return [c for i, c in enumerate(candidatos)
            if i < MAX_MESCLA or _HISTORICO.search(PurePath(c.arquivo.nome).stem)]


def mesclar(tipo: str, frames: list[pd.DataFrame]) -> pd.DataFrame:
    """Junta os arquivos de uma base (o 1º é o mais novo) pela chave; linhas sem chave
    — como as de fórmulas arrastadas até o fim da planilha — são descartadas."""
    from .util import achar_coluna, texto

    rx = MESCLAR.get(tipo)
    partes, chave_col = [], None
    for df in frames:
        col = achar_coluna(df.columns, rx) if rx else None
        if col:
            chave_col = chave_col or col
            df = df.copy()
            df[col] = texto(df[col])
            df = df[df[col] != ""].rename(columns={col: chave_col})
        partes.append(df)
    junto = partes[0] if len(partes) == 1 else pd.concat(partes, ignore_index=True, sort=False)
    if chave_col:
        junto = junto.drop_duplicates(chave_col, keep="first")
    return junto.reset_index(drop=True)


def identificar(colunas) -> str | None:
    ks = {chave(c) for c in colunas}

    def tem(rx):
        return any(re.search(rx, k) for k in ks)

    if tem(_COLUNAS_DO_APP):
        return None
    # Gerenciador de AF: ações (nº da análise + nº da ação + ação) e as análises de falha
    if tem(r"^N DA ANALISE$") and tem(r"^N DA ACAO$") and tem(r"^ACAO$"):
        return ACAO_AF
    if tem(r"^N DA ANALISE$") and tem(r"^STATUS DA ANALISE$") and tem(r"^DATA DA FALHA$"):
        return AF
    # IW47: confirmações (horas apontadas por pessoa). Antes do IW38, que também tem Ordem + Status.
    if tem(r"^TRABALHO REAL$") and tem(r"^ORDEM$") and tem(r"^N PESSOAL$|NUMERO PESSOAL"):
        return CONF
    # Gestão de HH: cadastro das pessoas da manutenção (nº pessoal, nome, cargo, centro de trabalho…)
    if tem(r"^N PESSOAL$|NUMERO PESSOAL|MATRICULA") and tem(r"^NOME") and tem(r"^CARGO$|^FUNCAO$") \
            and not tem(r"TRABALHO REAL"):
        return EQUIPE
    # IW38 com operações (IW38OP): uma linha por operação, com o trabalho planejado
    if tem(r"^ORDEM$") and tem(r"^OPERACAO$") and tem(r"^TRABALHO$|DURACAO NORMAL"):
        return OPER
    # IW28: notas de manutenção
    if tem(r"^NOTA$") and tem(r"TIPO DE NOTA|DATA DA NOTA"):
        return NOTAS
    # IH08: cadastro de equipamentos (sem ordem, nota, plano ou material)
    if tem(r"^EQUIPAMENTO$") and tem(r"DENOMINACAO") and not tem(r"^ORDEM$|^NOTA$|PLANO|^MATERIAL$|DATA BASE"):
        return EQUIP
    # IP19 (programação dos planos): plano + data planejada/de chamada. Vem antes do IW38
    # porque a IP19 também pode trazer a coluna Ordem; o IW38 tem "Data-base do início".
    if tem(r"^PLANO|PLANO (DE )?MANUT") and tem(r"DATA PLAN|DT PLAN|DATA (DE )?CHAMADA|DT CHAM") \
            and not tem(r"DATA BASE DO INICIO"):
        return IP19

    if tem(r"^ORDEM$") and (tem(r"STATUS USUARIO") or tem(r"STATUS DO SISTEMA") or tem(r"TIPO DE ORDEM")):
        return IW38
    if (tem(r"N(UMERO)? DA REQ") or tem(r"REQUISI")) and tem(r"^TOTAL$"):
        return REQ
    if tem(r"^MATERIAL$|^CODIGO") and tem(r"UTILIZACAO LIVRE|UTILIZ LIVRE|LIVRE UTILIZ|ESTOQUE|SALDO"):
        return MB52
    return None

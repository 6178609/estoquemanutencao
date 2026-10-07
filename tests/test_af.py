"""Gerenciador de AF: leitura das abas, situação das AFs e das ações e indicadores corporativos.
Formato reproduzido em miniatura, com dados fictícios."""

import io
from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook

from central import af, leitura
from central.leitura import ACAO_AF, AF, EQUIP

HOJE = pd.Timestamp("2026-06-30")
D = datetime

COLS_AF = ["Nº da análise", "Tipo de Análise", "Área", "Área da Falha", "Cod SAP", "Equipamento", "Conjunto",
           "Subconjunto", "Componente", "Pilar direcionador", "Causa Raiz", "Especial. da Causa",
           "Criticidade Equipamento", "Nº O.S Corretiva", "Resumo da Falha", "Duração da Parada",
           "Supervisor Responsável", "Turno", "Status da Análise", "Data da Falha", "Data Limite p/ Análise", "Ajuste",
           "Data da Análise", "Status das ações", "Link da AF", "Observações", "Id da Planilha", "Número"]
COLS_ACAO = ["Nº da Análise", "Nº da Ação", "Tipo de Análise", "Localização Macro", "Área", "Área da Falha", "Cod SAP",
             "Equipamento", "Especialidade de Correção", "O.S de Correção / Plano Sistêmico / Lista", "Link", "Ação",
             "Causa Raiz", "TTR", "Padronização", "Expansão", "Supervisor Responsável", "Responsável", "Status",
             "Data de Início", "Data Limite p/ Ação", "Data Realizada", "Ordem da Quebra", "Observações"]


def _af(n, crit, status, falha, limite, analise, sup="BETO", ajuste=None, sap=41000001.0):
    return [n, "GATILHO", "INJ. EXPANSIVOS", "ALPA-A026-2-3-2", sap, "INJETORA 01", "***", None, None, "mp",
            "MANUTENÇÃO INSUFICIENTE", "Mecânica", crit, n, f"AF - {n} - FALHA", "1.5", sup, "1º", status, falha,
            limite, ajuste, analise, 0.5, "https://exemplo/af", None, "/MP/", n]


def afs_cru():
    return pd.DataFrame([
        _af(32000001, "A", "AP", D(2026, 6, 1), D(2026, 6, 8), D(2026, 6, 5), sup="Diego Santos"),   # no prazo
        _af(32000002, "A", "AFP", D(2026, 6, 2), D(2026, 6, 9), D(2026, 6, 12), sup="DIEGO SANTOS "),  # fora
        _af(32000003, "A", None, D(2026, 6, 20), D(2026, 6, 27), None),                              # atrasada
        _af(32000004, "B", "R", D(2026, 6, 26), D(2026, 7, 3), None),                                # risco
        _af(32000005, "C", "EA", D(2026, 6, 27), D(2026, 7, 4), None),                               # andamento
        _af(32000006, "A", "AP", D(2026, 6, 3), D(2026, 6, 10), D(2026, 6, 4), ajuste="Cancelada"),  # cancelada
        _af(32000007, "C", None, D(2026, 6, 4), D(2026, 6, 11), D(2026, 6, 10)),                     # sem código
    ], columns=COLS_AF)


def _acao(n_af, n, acao, status, ini, limite, feita, resp="jeferson silva", cr=True, ttr=False, pad="x", exp="########"):
    return [n_af, n, "GATILHO", None, None, None, None, None, None, None, None, acao, cr, ttr, pad, exp, "BETO",
            resp, status, ini, limite, feita, n_af, None]


def acoes_cru():
    return pd.DataFrame([
        _acao(32000001, 1, "Trocar rolamento", "AP", D(2026, 6, 5), D(2026, 6, 15), D(2026, 6, 10)),
        _acao(32000001, 2, "Incluir no plano", "AFP", D(2026, 6, 5), D(2026, 6, 15), D(2026, 6, 20), cr=False),
        _acao(32000002, 1, "Revisar procedimento", None, D(2026, 6, 12), D(2026, 6, 20), None),     # atrasada
        _acao(32000002, 2, "Treinar equipe", None, D(2026, 6, 12), D(2026, 7, 3), None),            # vence em 7 d
        _acao(32000003, 1, "Analisar desenho", "EA", D(2026, 6, 25), D(2026, 7, 30), None),         # andamento
        _acao(32000003, 2, "Comprar peça", None, D(2026, 6, 25), D(2026, 8, 30), None),             # não iniciada
        _acao(32000004, 1, "Cancelada", "CANCELADA", D(2026, 6, 26), D(2026, 6, 28), D(2026, 6, 28)),
        _acao(99999999, 1, "AF antiga fora da aba", None, D(2026, 6, 1), D(2026, 6, 5), D(2026, 6, 4)),
        _acao(32000005, 1, None, None, D(2026, 6, 1), D(2026, 6, 5), None),                         # sem ação: fora
    ], columns=COLS_ACAO)


def gerenciador() -> bytes:
    """Pasta de trabalho no formato do Gerenciador de AF, com as abas ocultas que confundiriam a leitura."""
    wb = Workbook()
    g = wb.active
    g.title = "Gargalos"
    g.sheet_state = "hidden"
    g.append(["RELAÇÃO DOS EQUIPAMENTOS GARGALOS", None, None, "Equipamento", "Denominação", "Código ABC"])
    g.append([None, None, None, "41000002", "MISTURADOR", "A"])
    b = wb.create_sheet("Análise de Falha (Backup)")
    b.sheet_state = "hidden"
    b.append(COLS_AF)
    b.append(_af(1, "A", "AP", D(2023, 7, 4), D(2023, 7, 14), D(2023, 7, 10)))
    for nome, cols, linhas in [("Análise de Falha", COLS_AF, afs_cru().values.tolist()),
                               ("Plano de ação AF", COLS_ACAO, acoes_cru().values.tolist())]:
        ws = wb.create_sheet(nome)
        ws.append(cols)
        for linha in linhas:
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in linha])
    e = wb.create_sheet("E-mails")
    e.append(["Carimbo de data/hora", "Endereço de e-mail", "Area", "Cargo"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_identifica_so_as_abas_visiveis_de_af_e_acoes():
    achados = leitura.sondar("Gerenciador de AF - F26 2026.xlsx", gerenciador())
    assert achados == [(AF, "Análise de Falha", 0), (ACAO_AF, "Plano de ação AF", 0)]  # backup e Gargalos fora
    assert leitura.identificar(["Equipamento", "Denominação", "Código ABC"]) == EQUIP   # o IH08 continua IH08


def test_situacao_das_afs():
    a = af.preparar_afs(afs_cru(), HOJE).set_index("Nº AF")
    assert a.loc["32000001", "Situação"] == af.AF_NO_PRAZO
    assert a.loc["32000002", "Situação"] == af.AF_FORA_PRAZO and a.loc["32000002", "Dias de atraso"] == 3
    assert a.loc["32000003", "Situação"] == af.AF_ATRASADA and a.loc["32000003", "Dias de atraso"] == 3
    assert a.loc["32000004", "Situação"] == af.AF_RISCO
    assert a.loc["32000005", "Situação"] == af.AF_ANDAMENTO
    assert a.loc["32000006", "Situação"] == af.AF_CANCELADA
    assert a.loc["32000007", "Situação"] == af.AF_NO_PRAZO                 # sem código: decide pela data
    assert a.loc["32000001", "Supervisor"] == a.loc["32000002", "Supervisor"] == "DIEGO SANTOS"
    assert a.loc["32000001", "Pilar"] == "MP" and a.loc["32000001", "Conjunto"] == ""  # "***" vira vazio
    assert a.loc["32000001", "% ações concluídas"] == 50 and a.loc["32000001", "Dias para análise"] == 4
    assert a.loc["32000001", "Código SAP"] == "41000001" and a.loc["32000001", "Duração da parada (h)"] == 1.5


def test_situacao_das_acoes():
    a = af.preparar_afs(afs_cru(), HOJE)
    c = af.preparar_acoes(acoes_cru(), HOJE, a).set_index("ID")
    assert len(c) == 8                                                     # linha sem texto da ação fica fora
    assert c.loc["32000001-1", "Situação"] == af.AC_NO_PRAZO
    assert c.loc["32000001-2", "Situação"] == af.AC_FORA_PRAZO and c.loc["32000001-2", "Dias de atraso"] == 5
    assert c.loc["32000002-1", "Situação"] == af.AC_ATRASADA and c.loc["32000002-1", "Dias de atraso"] == 10
    assert c.loc["32000002-2", "Situação"] == af.AC_VENCE
    assert c.loc["32000003-1", "Situação"] == af.AC_ANDAMENTO
    assert c.loc["32000003-2", "Situação"] == af.AC_NAO_INICIADA
    assert c.loc["32000004-1", "Situação"] == af.AC_CANCELADA
    assert c.loc["32000001-1", "Responsável"] == "JEFERSON SILVA"
    # marcações: True / "x" = marcado; False / "########" = não
    assert c.loc["32000001-1", "Tipo da ação"] == "Causa raiz, Padronização"
    assert not c.loc["32000001-1", "Expansão"] and not c.loc["32000001-1", "TTR"]
    # dados da AF completam a ação
    assert c.loc["32000002-1", "Criticidade"] == "A" and c.loc["32000002-1", "Equipamento"] == "INJETORA 01"
    assert c.loc["32000002-1", "AF cadastrada"] and not c.loc["99999999-1", "AF cadastrada"]
    assert c.loc["32000001-1", "Dias para concluir"] == 5


def test_indicadores_corporativos():
    a = af.preparar_afs(afs_cru(), HOJE)
    c = af.preparar_acoes(acoes_cru(), HOJE, a)
    r = af.indicadores(a, c, date(2026, 6, 1), date(2026, 6, 30), HOJE)
    assert r["af_geradas"] == 6 and r["af_geradas_a"] == 3                 # cancelada não conta
    # taxa de quebra crítica A: AFs A com falha no período analisadas (1, 2) ÷ AFs A (1, 2, 3)
    assert r["taxa_quebra_a"] == pytest.approx(200 / 3) and r["af_a_no_prazo"] == pytest.approx(100 / 3)
    # execução: data limite no período até hoje → 1, 2, 3, 7; analisadas 1, 2, 7; no prazo 1, 7
    assert r["af_execucao"] == pytest.approx(75) and r["af_no_prazo"] == pytest.approx(50)
    assert r["af_atrasadas"] == 1 and r["af_abertas"] == 3
    # ações com data limite em junho até hoje (canceladas fora): 1-1, 1-2, 2-1, 9999-1 → 3 realizadas, 2 no prazo
    assert r["acoes_execucao"] == pytest.approx(75) and r["acoes_no_prazo"] == pytest.approx(50)
    assert r["acoes_atrasadas"] == 1 and r["acoes_abertas"] == 4
    assert r["quebras_a_iw28"] is None


def test_semanal_e_mensal():
    a = af.preparar_afs(afs_cru(), HOJE)
    c = af.preparar_acoes(acoes_cru(), HOJE, a)
    s = af.semanal(a, c, date(2026, 6, 1), date(2026, 6, 30), HOJE)
    sem = s[s["Semana"] == pd.Timestamp("2026-06-08")].set_index("Indicador")
    # semana de 08/06: AFs com limite 08, 09 e 11/06 (1, 2, 7) → todas analisadas
    assert sem.loc["Execução de AFs", "Gerado"] == 3 and sem.loc["Execução de AFs", "Percentual"] == 100
    assert set(s["Indicador"]) == {"Taxa de análise de quebra crítica (A)", "Execução de AFs", "Ações de AFs"}
    m = af.mensal(a, c, date(2026, 5, 1), date(2026, 6, 30), HOJE)
    assert list(m["Mês"]) == [pd.Timestamp("2026-05-01"), pd.Timestamp("2026-06-01")]
    assert m.loc[1, "af_geradas"] == 6 and m.loc[0, "af_geradas"] == 0


def test_quebras_a_da_iw28_como_referencia():
    notas = pd.DataFrame({"Nota": ["1", "2", "3"], "Tipo de nota": ["Y1", "Y1", "Y2"], "Ordem": ["9", "8", "7"],
                          "Código ABC": ["A", "B", "A"],
                          "Data": [pd.Timestamp("2026-06-02")] * 3, "Equip. (chave)": ["E1", "E2", "E3"]})
    a = af.preparar_afs(afs_cru(), HOJE)
    c = af.preparar_acoes(acoes_cru(), HOJE, a)
    assert af.indicadores(a, c, date(2026, 6, 1), date(2026, 6, 30), HOJE, notas)["quebras_a_iw28"] == 1
    assert list(af.acoes_da_af(c, "32000001")["Nº ação"]) == ["1", "2"]

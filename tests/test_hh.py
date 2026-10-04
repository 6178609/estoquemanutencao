"""IW47 (apontamentos de horas) e planilha de Gestão de HH (só a lista de pessoas).
Formatos reproduzidos em miniatura, com dados fictícios."""

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook

from central import bases, leitura
from central.leitura import CONF, EQUIPE, IW38

COLS_IW47 = ["Nº pessoal", "Ordem", "Status do sistema", "Operação", "Suboperação", "Trabalho real",
             "Unid.trabalho (real)", "Centro trab.(real)", "Tp.atividade (real)", "Tipo ativid.(plan.)",
             "Data lançamento", "Criado em"]


def iw47():
    return pd.DataFrame([
        [1001, 32000001, "CONF ENTE", 10, None, 90, "MIN", "MEC_COMP", "MECA", "MECA", datetime(2026, 6, 8), datetime(2026, 6, 8)],
        [1001, 32000001, "CONF ENTE", 20, None, 30, "MIN", "MEC_COMP", None, "ELETRI", datetime(2026, 6, 8), datetime(2026, 6, 8)],
        [1002, 32000002, "CONF ENTE", 10, None, 600, "MIN", "ELE_COMP", "ELETRI", "ELETRI", datetime(2026, 6, 9), datetime(2026, 6, 9)],
        [1002, 32000002, "CONF ENTE", 10, None, -600, "MIN", "ELE_COMP", "ELETRI", "ELETRI", datetime(2026, 6, 9), datetime(2026, 6, 9)],
        [1003, 32000003, "CONF ENTE", 10, None, 60, "MIN", "FABPILOT", "MECA", "MECA", datetime(2026, 6, 9), datetime(2026, 6, 9)],
    ], columns=COLS_IW47)


def _xlsx(df):
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    return buf.getvalue()


def gestao_hh() -> bytes:
    """Pasta de trabalho no formato da Gestão de HH: lista de pessoas (cabeçalho na linha 5), aba de apoio com
    os tipos de OM e cópias antigas da IW47 e do IW38 (que não podem virar base)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BD 2"
    ws.append(["DES_POSTO", "TECNICO", None, "TIPO", "TIPO DE OM"])
    ws.append(["MECANICO", "MECANICO", None, "YM11", "Corretiva Emergencial"])
    ws.append(["ELETRICISTA", "ELETRICISTA", None, "YM15", "Preventiva"])
    g = wb.create_sheet("Gestão de HH")
    g.append([None, 1, 2, 3])
    g.append([])
    g.append([])
    g.append([])
    g.append(["Nº", "Nº PESSOAL", "NOME", "CARGO", "CENTRO TRABALHO", "AREA", "TURMA", datetime(2026, 1, 1)])
    g.append([1, 1001, "ANA SOUZA", "MECANICO ESPECIALIZADO", "MEC_COMP", "COMPONENTES", "1ª", 8])
    g.append([2, 1002, "BRUNO LIMA", "ELETRICISTA LIDER", "ELE_COMP", "COMPONENTES", "2ª", 8])
    g.append([3, 1004, "CARLA DIAS", "GPM", "GPM", "PLANEJADA", "#REF!", 8])
    g.append([4, None, "#REF!", "#REF!", None, None, None, None])
    c = wb.create_sheet("IW47")
    c.append(["ÁREA", "Ordem", "Operação", "Centro trab.(real)", "Trabalho real", "Nº pessoal", "Nome do empregado"])
    c.append(["X", 31968702, 100, "MEC_MIS", 10, 1001, "ANA SOUZA"])
    x = wb.create_sheet("DadosExtras (IW38)")
    x.append(["Ordem", "Tipo de ordem", "Tipo ativid.PM", "Denominação", "Plano manut."])
    x.append([32004539, "YM11", 111, "SILK", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_iw47_nao_e_confundida_com_iw38():
    assert leitura.identificar(COLS_IW47) == CONF
    assert leitura.sondar("IW47.XLSX", _xlsx(iw47())) == [(CONF, "Sheet1", 0)]
    # o IW38 continua sendo IW38
    assert leitura.identificar(["Ordem", "Tipo de ordem", "Status do sistema", "Texto breve"]) == IW38


def test_preparar_iw47():
    df = bases.preparar_confirmacoes(iw47())
    assert set(df["Ordem"]) == {"32000001", "32000002"}            # FABPILOT fora
    assert df["Horas"].sum() == 2.0                                  # 90 + 30 min; 600 − 600 (estorno) = 0
    assert df.loc[df["Operação"] == "20", "Atividade"].iat[0] == "ELETRI"  # sem a real, vale a planejada
    assert df["Nº pessoal"].iat[0] == "1001" and df["Data"].iat[0] == pd.Timestamp("2026-06-08")
    assert len(bases.preparar_confirmacoes(iw47(), excluir_piloto_matriz=False)) == 5


def test_gestao_de_hh_so_vale_a_lista_de_pessoas():
    achados = leitura.sondar("1-Gestão de Hh F26 (1).xlsx", gestao_hh())
    assert achados == [(EQUIPE, "Gestão de HH", 4)]                  # cópias de IW47/IW38 ignoradas
    cru = leitura.ler_arquivo("g.xlsx", gestao_hh(), "Gestão de HH", 4)
    eq = bases.preparar_equipe(cru)
    assert list(eq["Nº pessoal"]) == ["1001", "1002", "1004"]        # linha #REF! fora
    assert list(eq["Especialidade"]) == ["Mecânica", "Elétrica", "Planejamento (GPM)"]
    assert eq.loc[eq["Nº pessoal"] == "1004", "Turma"].iat[0] == ""   # #REF! vira vazio
    assert eq["Área"].iat[0] == "COMPONENTES"


def test_tipos_de_ordem_da_aba_de_apoio(tmp_path, monkeypatch):
    (tmp_path / "Gestão de HH.xlsx").write_bytes(gestao_hh())
    monkeypatch.setenv("CENTRAL_PASTAS", str(tmp_path))
    monkeypatch.setenv("CENTRAL_PASTA_APP", str(tmp_path / "app"))
    bases._fonte_cacheada.clear()
    bases.recarregar()
    try:
        assert bases.equipe().df is not None and len(bases.equipe().df) == 3
        assert bases.tipos_de_ordem_padrao() == {"YM11": "Corretiva Emergencial", "YM15": "Preventiva"}
        assert bases.iw38().df is None and bases.confirmacoes().df is None  # cópias da planilha não viram base
    finally:
        bases._fonte_cacheada.clear()
        bases.recarregar()
